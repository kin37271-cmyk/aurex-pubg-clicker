# excel_export.py
import os
import io
import aiosqlite
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from database import DB_NAME

async def generate_users_excel() -> str:
    """Barcha foydalanuvchilar ma'lumotlarini Excel (.xlsx) faylga chiqarish"""
    file_path = "foydalanuvchilar.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Foydalanuvchilar"

    # Headerlar
    headers = [
        "Telegram ID", "Ism Familiya", "Username", "Balans (🪙)", 
        "Jami Ishlangan (🪙)", "Energiya", "Max Quvvat", "Multitap Lvl", 
        "Quvvat Lvl", "Regen Lvl", "Auto-Bot Lvl", "Do'stlari Soni", 
        "Kim Taklif Qilgan ID", "Holat (Ban)", "Ro'yxatdan O'tgan Sana"
    ]
    ws.append(headers)

    # Header dizayni (To'q ko'k rang, oq qalin yozuv)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Ma'lumotlarni bazadan olish
    async with aiosqlite.connect(DB_NAME) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute("SELECT * FROM users ORDER BY created_at DESC") as cursor:
            rows = await cursor.fetchall()

    alt_fill = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")

    for row_idx, r in enumerate(rows, start=2):
        ban_status = "🔴 Bloklangan" if r["is_banned"] else "🟢 Faol"
        username_val = f"@{r['username']}" if r["username"] else "Mavjud emas"
        ref_val = r["referred_by"] if r["referred_by"] else "-"
        
        row_data = [
            str(r["user_id"]),
            r["full_name"] or "",
            username_val,
            round(float(r["balance"]), 2),
            round(float(r["total_earned"]), 2),
            r["energy"],
            r["max_energy"],
            r["multitap_level"],
            r["energy_level"],
            r["regen_level"],
            r["autobot_level"],
            r["referral_count"],
            str(ref_val),
            ban_status,
            str(r["created_at"] or "")
        ]
        ws.append(row_data)

        # Qatorlar bezagi
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if is_even:
                cell.fill = alt_fill
            if col_idx in [1, 13, 15]:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in [4, 5, 6, 7, 8, 9, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="right")

    # Ustun kengliklarini avtomatik to'g'rilash
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(file_path)
    return file_path

async def generate_withdrawals_excel() -> str:
    """Barcha pul yechish zayafkalarini Excel (.xlsx) faylga chiqarish"""
    file_path = "zayafkalar.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Zayafkalar"

    headers = [
        "Zayafka ID", "Foydalanuvchi ID", "Ismi", "Turi (Karta/PUBG)", 
        "Yechilgan Tanga (🪙)", "To'lanadigan Summa / UC", "Karta / PUBG ID", 
        "Karta Egasi / Nickname", "Holati", "So'rov Sanasi", 
        "Hal Qilingan Sana", "Admin Izohi"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E7145", end_color="1E7145", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    async with aiosqlite.connect(DB_NAME) as db:
        db.row_factory = aiosqlite.Row
        query = """
        SELECT w.*, u.full_name 
        FROM withdrawals w 
        LEFT JOIN users u ON w.user_id = u.user_id 
        ORDER BY w.id DESC
        """
        async with db.execute(query) as cursor:
            rows = await cursor.fetchall()

    status_map = {
        "PENDING": "⏳ Kutilmoqda",
        "APPROVED": "✅ To'landi",
        "REJECTED": "❌ Rad etildi"
    }

    alt_fill = PatternFill(start_color="F4FAF6", end_color="F4FAF6", fill_type="solid")

    for row_idx, r in enumerate(rows, start=2):
        target_unit = "so'm" if r["type"] == "CARD" else "UC"
        status_text = status_map.get(r["status"], r["status"])
        
        row_data = [
            f"#{r['id']}",
            str(r["user_id"]),
            r["full_name"] or "Noma'lum",
            r["type"],
            round(float(r["amount_coins"]), 2),
            f"{float(r['amount_target']):,.0f} {target_unit}",
            str(r["target_value"] or ""),
            str(r["target_details"] or ""),
            status_text,
            str(r["created_at"] or ""),
            str(r["resolved_at"] or "-"),
            str(r["admin_note"] or "-")
        ]
        ws.append(row_data)

        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if is_even:
                cell.fill = alt_fill
            if col_idx in [1, 2, 4, 9, 10, 11]:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in [5, 6]:
                cell.alignment = Alignment(horizontal="right")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(file_path)
    return file_path
